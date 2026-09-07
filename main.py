from fastapi import FastAPI, UploadFile, File, Form, HTTPException, BackgroundTasks, Request
from fastapi.responses import FileResponse, JSONResponse
from fastapi.concurrency import run_in_threadpool
from pdf2docx import Converter
import pdfplumber
from openpyxl import Workbook
from pptx import Presentation
import fitz # PyMuPDF
import tempfile
import os
from pathlib import Path
import re
import secrets
import shutil
import socket
import subprocess
import threading
import time
import zipfile
import edge_tts
from pydantic import BaseModel, Field
import imageio_ffmpeg
from unoserver.client import UnoClient

# LibreOffice (word/excel/ppt <-> pdf) and Ghostscript (compress-pdf, thumbnails) run here
# rather than in the Vercel Next.js app because neither fits a serverless function: Vercel
# caps deployment size at 250MB and LibreOffice alone is 300MB-1GB+, so it can never be
# bundled there regardless of install method. This host is a normal Linux container/VM, so
# both are installed as regular system packages - see the Dockerfile.
SOFFICE_BIN = os.environ.get("SOFFICE_PATH", "soffice")
GHOSTSCRIPT_BIN = os.environ.get("GHOSTSCRIPT_PATH", "gs")

# word/excel/ppt -> pdf normally spawns a brand-new LibreOffice process per request (see
# _convert_with_soffice_spawn below), and LibreOffice's own startup - initializing its UNO
# service manager, fonts, config - is a fixed ~10-15s cost regardless of document size. Setting
# this keeps one LibreOffice instance running persistently instead (via the `unoserver` project)
# and routes conversions through it, paying that startup cost once instead of on every request -
# see convert_via_uno_listener. Falls back automatically to the slower per-request spawn if the
# listener can't start, or (per-file-extension, see the circuit breaker below) if it keeps
# crashing on a given format, so this is safe to leave on even where the listener doesn't work.
USE_PERSISTENT_LIBREOFFICE = os.environ.get("USE_PERSISTENT_LIBREOFFICE", "true").lower() not in ("false", "0", "no")
# Must be a Python interpreter that has LibreOffice's `uno` bridge module importable - that's
# never this service's own venv (installing `uno` there isn't a normal pip package). On the
# Linux/Docker deployment that's the system python3 after `apt-get install python3-uno` (see
# Dockerfile); on Windows it's LibreOffice's bundled program\python.exe. Only the *listener*
# process (spawned with this interpreter) needs the uno bridge - this service talks to it as a
# plain network client via the `unoserver` pip package (see requirements.txt), which needs no
# uno import of its own.
UNOSERVER_PYTHON = os.environ.get("UNOSERVER_PYTHON", "python3")
UNO_LISTENER_HOST = os.environ.get("UNO_LISTENER_HOST", "127.0.0.1")
UNO_LISTENER_PORT = os.environ.get("UNO_LISTENER_PORT", "2003")
UNO_INTERNAL_PORT = os.environ.get("UNO_INTERNAL_PORT", "2002")

# Automatically append the bundled ffmpeg to the system PATH
os.environ["PATH"] += os.pathsep + os.path.dirname(imageio_ffmpeg.get_ffmpeg_exe())

import faster_whisper
# Load model once at startup to avoid loading on every request
whisper_model = faster_whisper.WhisperModel("base", device="cpu", compute_type="int8")

class TTSRequest(BaseModel):
    # Mirrors the 5000-char cap enforced client-side (TextToSpeechZone.tsx) and
    # server-side (the Next.js route) - kept here too so a request that reaches
    # this service directly can't force an oversized edge-tts synthesis job.
    text: str = Field(..., max_length=5000)
    voice: str = "en-US-JennyNeural"
    speed: str = "+0%"

app = FastAPI(title="Convertify Python Microservice")

# Shared secret with the Next.js app (lib/pythonService.ts) so /convert/* can't be
# hit for free by anyone who can route to this service directly, bypassing the
# rate limiting that otherwise only runs in the Next.js route handlers. Optional -
# if unset, the check is skipped so local dev doesn't need it configured - but the
# Next.js side must be given the same value once this is exposed publicly.
PYTHON_SERVICE_SECRET = os.environ.get("PYTHON_SERVICE_SECRET")
if not PYTHON_SERVICE_SECRET:
    print(
        "[startup] WARNING: PYTHON_SERVICE_SECRET is not set - /convert/* endpoints "
        "are reachable by anyone who can route to this service. Set it here and as "
        "PYTHON_SERVICE_SECRET in the Next.js app before exposing this service publicly."
    )


@app.middleware("http")
async def verify_internal_secret(request: Request, call_next):
    if PYTHON_SERVICE_SECRET and request.url.path.startswith("/convert/"):
        expected = f"Bearer {PYTHON_SERVICE_SECRET}"
        provided = request.headers.get("authorization", "")
        if not secrets.compare_digest(provided, expected):
            return JSONResponse(status_code=401, content={"detail": "Unauthorized"})
    return await call_next(request)


@app.get("/")
def read_root():
    return {"message": "Python Microservice is running!"}


@app.on_event("startup")
async def _warm_up_uno_listener():
    """Best-effort: start the persistent LibreOffice listener at boot so the first real
    word/excel/ppt<->pdf request doesn't have to pay its ~10-15s startup cost itself. If this
    fails (e.g. UNOSERVER_PYTHON not configured for this host), convert_with_soffice() falls
    back to the per-request spawn transparently - so this is never fatal to startup."""
    if USE_PERSISTENT_LIBREOFFICE:
        await run_in_threadpool(_ensure_uno_listener)


@app.on_event("shutdown")
async def _shut_down_uno_listener():
    with _uno_listener_state_lock:
        _stop_uno_listener_locked()

def remove_files(paths):
    for path in paths:
        try:
            if os.path.exists(path):
                os.remove(path)
        except Exception as e:
            print(f"Error removing {path}: {e}")


class InvalidFileError(Exception):
    """Raised when the uploaded file itself is the problem - corrupted, empty,
    password-protected, or not actually the format its extension claims - as
    opposed to an unexpected server/infrastructure failure. Endpoints turn this
    into a 422 with the message shown directly to the user, so messages here
    must always be safe to show as-is: no filesystem paths, no raw library
    exception text. Kept distinct from a plain 500 so the Next.js caller (and
    ultimately the end user) can tell "your file has a problem, fix it and
    re-upload" apart from "something is wrong on our end, retrying may help" -
    collapsing both into one generic error was actively misleading users
    whose file was the actual issue (e.g. password-protected or corrupted)."""
    pass


# OOXML (docx/xlsx/pptx) files are zip archives with a predictable member for
# their main content part; legacy (doc/xls/ppt) files are OLE2 compound files
# with a fixed 8-byte magic number. Checking these before handing a file to
# LibreOffice matters because --convert-to auto-detects the import filter and
# will happily import a garbage/plain-text file as a *text document*, silently
# producing a "successful" PDF that just contains the raw garbage as text -
# instead of failing, which is what a corrupted or mislabeled upload should do.
_OOXML_MAIN_PART = {
    ".docx": "word/document.xml",
    ".xlsx": "xl/workbook.xml",
    ".pptx": "ppt/presentation.xml",
}
_LEGACY_OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_KIND_BY_EXT = {
    ".doc": "Word", ".docx": "Word",
    ".xls": "Excel", ".xlsx": "Excel",
    ".ppt": "PowerPoint", ".pptx": "PowerPoint",
}


def validate_office_file(input_path: str, ext: str) -> None:
    """Rejects a file that doesn't actually match the Office format its extension
    claims, before it ever reaches LibreOffice. See InvalidFileError and the
    module-level comment above for why this check exists."""
    ext = ext.lower()
    kind = _KIND_BY_EXT.get(ext, "document")
    bad_file_message = (
        f"This doesn't look like a valid {kind} file. It may be corrupted, empty, or not "
        f"actually a {kind} document - please check the file and try again."
    )

    if ext in _OOXML_MAIN_PART:
        try:
            with zipfile.ZipFile(input_path) as zf:
                names = zf.namelist()
                if "[Content_Types].xml" not in names or _OOXML_MAIN_PART[ext] not in names:
                    raise InvalidFileError(bad_file_message)
        except zipfile.BadZipFile:
            raise InvalidFileError(bad_file_message)
    elif ext in (".doc", ".xls", ".ppt"):
        with open(input_path, "rb") as f:
            header = f.read(8)
        if header != _LEGACY_OLE_MAGIC:
            raise InvalidFileError(bad_file_message)


def open_pdf_or_raise(pdf_path: str) -> fitz.Document:
    """Opens a PDF for the pdf-to-ppt pipeline with a clear, user-safe
    InvalidFileError for the two most common bad-input cases - unreadable/corrupt
    file and password-protected file - instead of letting PyMuPDF's low-level
    errors (e.g. "Failed to open file 'C:\\Users\\...\\tmp78ysehkd.pdf' as type
    pdf" or "document closed or encrypted") reach the client. Those messages are
    both cryptic to an end user and, in the corrupt-file case, leak the server's
    internal temp file path."""
    try:
        doc = fitz.open(pdf_path)
    except Exception:
        raise InvalidFileError(
            "This doesn't look like a valid PDF file. It may be corrupted or not actually "
            "a PDF - please check the file and try again."
        )
    if doc.needs_pass:
        doc.close()
        raise InvalidFileError(
            "This PDF is password-protected. Please remove the password (e.g. via your PDF "
            "viewer's \"Print to PDF\" or a password-removal tool) and upload it again."
        )
    return doc


def _convert_with_soffice_spawn(input_path: str, output_dir: str, target_format: str = "pdf") -> str:
    """The original conversion path: spawns a brand-new LibreOffice process for this one
    conversion, with its own throwaway -env:UserInstallation profile dir so concurrent requests
    don't collide on LibreOffice's single-instance profile lock. Slower (~10-15s LibreOffice
    startup cost paid on every single call, regardless of document size) but has no shared,
    crashable state - used directly when USE_PERSISTENT_LIBREOFFICE is off, and as the automatic
    fallback from convert_with_soffice() when the persistent listener is unavailable or unreliable
    for a given file type."""
    profile_dir = tempfile.mkdtemp()
    try:
        # Path(...).as_uri() (not a manual f"file://{profile_dir}") because that manual form is
        # only a valid URI on POSIX - on Windows a raw path like "C:\Users\...\tmp" produces
        # "file://C:\Users\...\tmp", which isn't a well-formed file URI (drive letter without
        # the required extra slash, backslashes instead of forward slashes). LibreOffice then
        # fails to parse it and exits with code 1 and no stderr output at all, which is exactly
        # what running this in local Windows dev (rather than the Linux Docker/Railway host)
        # surfaces - as_uri() produces the correct form on both platforms.
        profile_uri = Path(profile_dir).as_uri()
        result = subprocess.run(
            [
                SOFFICE_BIN, "--headless", "--norestore",
                f"-env:UserInstallation={profile_uri}",
                "--convert-to", target_format, "--outdir", output_dir, input_path,
            ],
            capture_output=True,
            timeout=100,
        )
    finally:
        shutil.rmtree(profile_dir, ignore_errors=True)
    base_name = os.path.splitext(os.path.basename(input_path))[0]
    output_path = os.path.join(output_dir, f"{base_name}.{target_format}")
    if result.returncode != 0 or not os.path.exists(output_path):
        stderr = result.stderr.decode(errors="ignore") if result.stderr else "unknown error"
        raise RuntimeError(f"LibreOffice conversion failed: {stderr}")
    return output_path


# --- Persistent LibreOffice listener (see USE_PERSISTENT_LIBREOFFICE above) -----------------
#
# unoserver's own docs are explicit that it does *not* restart LibreOffice after a crash and
# expects whatever wraps it to handle that - the state and functions below are that wrapper:
# start/detect-death/restart the listener process, and a small per-file-extension circuit
# breaker so a format that reliably crashes the listener (observed in testing with some Impress
# ->PDF exports) doesn't end up slower than never having a fast path at all (fast attempt that
# crashes + listener restart + slow fallback, repeated on every request for that format).
_uno_listener_process = None
_uno_listener_profile_dir = None
_uno_listener_state_lock = threading.Lock()   # guards starting/stopping the listener process
_uno_conversion_lock = threading.Lock()       # serializes conversions - only one at a time can use the single listener

_UNO_FAILURE_THRESHOLD = 2
_UNO_BREAKER_COOLDOWN_SECONDS = 30 * 60
_uno_failure_counts: dict = {}   # file extension -> consecutive fast-path failure count
_uno_breaker_until: dict = {}    # file extension -> timestamp until which the fast path is skipped
_uno_breaker_lock = threading.Lock()


def _uno_listener_reachable() -> bool:
    try:
        with socket.create_connection((UNO_LISTENER_HOST, int(UNO_LISTENER_PORT)), timeout=1):
            return True
    except OSError:
        return False


def _stop_uno_listener_locked() -> None:
    """Caller must hold _uno_listener_state_lock."""
    global _uno_listener_process, _uno_listener_profile_dir
    if _uno_listener_process is not None:
        try:
            _uno_listener_process.terminate()
            _uno_listener_process.wait(timeout=5)
        except Exception:
            try:
                _uno_listener_process.kill()
            except Exception:
                pass
        _uno_listener_process = None
    if _uno_listener_profile_dir:
        shutil.rmtree(_uno_listener_profile_dir, ignore_errors=True)
        _uno_listener_profile_dir = None


def _start_uno_listener_locked() -> bool:
    """(Re)starts the persistent LibreOffice listener. Caller must hold _uno_listener_state_lock.
    Returns whether it came up and is reachable - false means the caller should fall back to
    _convert_with_soffice_spawn instead."""
    global _uno_listener_process, _uno_listener_profile_dir
    _stop_uno_listener_locked()

    _uno_listener_profile_dir = tempfile.mkdtemp(prefix="uno_listener_profile_")
    try:
        _uno_listener_process = subprocess.Popen(
            [
                UNOSERVER_PYTHON, "-m", "unoserver.server",
                "--interface", UNO_LISTENER_HOST,
                "--port", UNO_LISTENER_PORT,
                "--uno-port", UNO_INTERNAL_PORT,
                "--user-installation", _uno_listener_profile_dir,
                # Without this, a conversion that hangs (rather than cleanly crashing) would
                # block the shared listener - and every request waiting on _uno_conversion_lock
                # behind it - forever. Hitting this also makes unoserver exit, which the crash
                # handling in convert_via_uno_listener already restarts from.
                "--conversion-timeout", "60",
            ],
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
        )
    except Exception as e:
        print(f"[uno-listener] failed to spawn ({UNOSERVER_PYTHON} -m unoserver.server): {e}")
        _uno_listener_process = None
        return False

    deadline = time.time() + 25
    while time.time() < deadline:
        if _uno_listener_process.poll() is not None:
            print(f"[uno-listener] process exited during startup (code {_uno_listener_process.returncode})")
            return False
        if _uno_listener_reachable():
            print("[uno-listener] ready")
            return True
        time.sleep(0.5)
    print("[uno-listener] did not become reachable within 25s")
    _stop_uno_listener_locked()
    return False


def _ensure_uno_listener() -> bool:
    with _uno_listener_state_lock:
        if _uno_listener_process is not None and _uno_listener_process.poll() is None:
            return True
        return _start_uno_listener_locked()


def _uno_breaker_open(ext: str) -> bool:
    with _uno_breaker_lock:
        until = _uno_breaker_until.get(ext)
        return until is not None and time.time() < until


def _uno_record_success(ext: str) -> None:
    with _uno_breaker_lock:
        _uno_failure_counts[ext] = 0
        _uno_breaker_until.pop(ext, None)


def _uno_record_failure(ext: str) -> None:
    with _uno_breaker_lock:
        count = _uno_failure_counts.get(ext, 0) + 1
        _uno_failure_counts[ext] = count
        if count >= _UNO_FAILURE_THRESHOLD:
            _uno_breaker_until[ext] = time.time() + _UNO_BREAKER_COOLDOWN_SECONDS
            print(
                f"[uno-listener] {count} consecutive failures converting {ext} - disabling the "
                f"fast path for it for {_UNO_BREAKER_COOLDOWN_SECONDS // 60} min"
            )


def convert_via_uno_listener(input_path: str, output_dir: str, target_format: str) -> str:
    """Converts via the persistent listener instead of spawning a new LibreOffice process.
    Raises RuntimeError if the listener isn't usable right now - convert_with_soffice() catches
    that and falls back to _convert_with_soffice_spawn()."""
    if not _ensure_uno_listener():
        raise RuntimeError("uno listener unavailable")

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    output_path = os.path.join(output_dir, f"{base_name}.{target_format}")
    client = UnoClient(server=UNO_LISTENER_HOST, port=UNO_LISTENER_PORT)

    with _uno_conversion_lock:
        try:
            client.convert(inpath=input_path, outpath=output_path, convert_to=target_format)
        except Exception as e:
            # Observed in testing: a conversion that crashes LibreOffice takes the whole
            # listener down with it, and unoserver does not restart itself after that (by
            # design, per its own docs) - restart it now so the *next* request gets a fresh,
            # working listener instead of repeatedly hitting a dead one.
            with _uno_listener_state_lock:
                _start_uno_listener_locked()
            raise RuntimeError(f"uno listener conversion failed: {e}")

    if not os.path.exists(output_path):
        raise RuntimeError("uno listener reported success but produced no output file")
    return output_path


def convert_with_soffice(input_path: str, output_dir: str, target_format: str = "pdf") -> str:
    """Converts a document to the given format via LibreOffice (e.g. word/excel/ppt -> pdf).
    Tries the fast persistent-listener path first, and transparently falls back to the original
    per-request LibreOffice spawn if the listener is disabled (USE_PERSISTENT_LIBREOFFICE=false),
    unavailable, or has repeatedly failed to convert this file extension recently (the circuit
    breaker above)."""
    ext = os.path.splitext(input_path)[1].lower()

    if USE_PERSISTENT_LIBREOFFICE and not _uno_breaker_open(ext):
        try:
            result = convert_via_uno_listener(input_path, output_dir, target_format)
            _uno_record_success(ext)
            return result
        except Exception as e:
            print(f"[uno-listener] fast path failed for '{ext}', falling back to per-request soffice: {e}")
            _uno_record_failure(ext)

    return _convert_with_soffice_spawn(input_path, output_dir, target_format)


def prepare_excel_for_pdf(input_path: str, ext: str) -> None:
    """Force Calc's 'fit sheet to page width' plus a tight print area before handing off
    to LibreOffice. Headless --convert-to otherwise leaves pagination entirely to whatever
    the workbook happened to save - for most files exported from a web app or built from
    scraped/pasted data that's nothing, so a wide sheet gets sliced across several PDF
    pages mid-table instead of shrinking to fit one page wide, the way Excel's own
    print preview normally behaves."""
    if ext not in (".xlsx", ".xlsm"):
        return
    try:
        from openpyxl import load_workbook
        wb = load_workbook(input_path)
        for ws in wb.worksheets:
            if ws.max_row < 1 or ws.max_column < 1:
                continue
            ws.print_area = ws.dimensions
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            if ws.max_column > 8:
                ws.page_setup.orientation = "landscape"
        wb.save(input_path)
    except Exception as e:
        # Best-effort only - if the workbook can't be parsed (password-protected,
        # unusual dialect), fall back to LibreOffice's default conversion rather than
        # failing the whole request.
        print(f"[excel-to-pdf] print-setup pre-processing skipped: {e}")


def _office_to_pdf_sync(file_obj, ext: str, work_dir: str) -> str:
    """The actual blocking work for word/excel/ppt -> pdf: writes the upload to disk and
    runs it through LibreOffice. Must be called via run_in_threadpool - convert_with_soffice()
    can take up to its 100s subprocess timeout, and calling it directly on the event loop
    would freeze every other request this service is handling (every other conversion tool,
    for every other user) for that entire duration."""
    input_path = os.path.join(work_dir, f"input{ext}")
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    validate_office_file(input_path, ext)
    prepare_excel_for_pdf(input_path, ext)

    return convert_with_soffice(input_path, work_dir)


async def office_to_pdf_response(background_tasks: BackgroundTasks, file: UploadFile) -> FileResponse:
    """Shared handler for word/excel/ppt -> pdf: saves the upload, converts via LibreOffice,
    and returns the resulting PDF. Keeps the original extension so LibreOffice picks the
    right import filter (works for both legacy .doc/.xls/.ppt and modern .docx/.xlsx/.pptx)."""
    ext = os.path.splitext(file.filename or "")[1] or ".tmp"
    work_dir = tempfile.mkdtemp()

    try:
        output_path = await run_in_threadpool(_office_to_pdf_sync, file.file, ext, work_dir)

        background_tasks.add_task(shutil.rmtree, work_dir, True)
        return FileResponse(path=output_path, filename=f"converted_{file.filename}.pdf", media_type="application/pdf")
    except InvalidFileError as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        print(f"[office-to-pdf] conversion failed: {e}")
        raise HTTPException(
            status_code=500,
            detail="We couldn't convert this file due to an unexpected server error. Please try again shortly.",
        )

def _convert_pdf_to_word_sync(file_obj, temp_pdf_path: str, temp_docx_path: str) -> str:
    """All the blocking work for pdf-to-word: writing the upload to disk, the pdf2docx
    conversion, and the spacing-fix passes below (all synchronous, CPU/IO-bound). Must be
    called via run_in_threadpool - otherwise a large/complex PDF blocks every other request
    this service is handling (every other conversion tool, for every other user) for
    however long the conversion takes."""
    # Write the uploaded file to the temp PDF path
    with open(temp_pdf_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    from pdf2docx import Converter
    import docx
    import fitz

    def has_spacing_issue(docx_path):
        try:
            doc = docx.Document(docx_path)
            long_words = 0
            total_words = 0
            for para in doc.paragraphs:
                words = para.text.split()
                for w in words:
                    total_words += 1
                    if len(w) >= 15:
                        long_words += 1
            if total_words == 0: return False
            return (long_words / total_words) > 0.05
        except:
            return False

    def fix_glued_words_in_docx(docx_path):
        # pdf2docx occasionally drops the space between words (common on
        # Canva-exported PDFs, justified text, etc.), which has_spacing_issue()
        # detects. Rather than throwing away pdf2docx's layout (tables, images,
        # columns, fonts - all correct) and rebuilding the page from scratch,
        # try a much cheaper surgical fix first: re-segment only the glued
        # tokens in place with a dictionary-based word splitter, leaving every
        # other run and all formatting untouched. Falls back to the full
        # from-scratch rebuild only if this doesn't clear the issue.
        import wordninja

        def split_token(token):
            if not token.isalpha() or len(token) < 15:
                return token
            parts = wordninja.split(token)
            if len(parts) <= 1 or ''.join(parts) != token:
                return token
            return ' '.join(parts)

        def fix_paragraphs(paragraphs):
            for para in paragraphs:
                for run in para.runs:
                    if not run.text:
                        continue
                    words = run.text.split(' ')
                    fixed = [split_token(w) for w in words]
                    if fixed != words:
                        run.text = ' '.join(fixed)

        try:
            doc = docx.Document(docx_path)
            fix_paragraphs(doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        fix_paragraphs(cell.paragraphs)
            doc.save(docx_path)
            return not has_spacing_issue(docx_path)
        except:
            return False

    def convert_with_correct_spacing(pdf_path, docx_path):
        import re
        from io import BytesIO
        from collections import Counter
        from docx.shared import Pt, Inches
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        pdf = fitz.open(pdf_path)
        out_doc = docx.Document()
        section = out_doc.sections[0]
        usable_width_in = section.page_width.inches - section.left_margin.inches - section.right_margin.inches

        def clean_font_name(f_str):
            name = f_str.split('+')[-1]
            name = re.sub(r'-(Bold|Italic|Regular|Medium|SemiBold|Light).*', '', name, flags=re.IGNORECASE)
            name = re.sub(r'(MT|PSMT|MS|PS)$', '', name, flags=re.IGNORECASE)
            name = re.sub(r'([a-z])([A-Z])', r'\1 \2', name)
            return name.strip()

        def add_page_number_field(paragraph):
            run = paragraph.add_run()
            begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
            instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = "PAGE"
            end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
            run._r.append(begin); run._r.append(instr); run._r.append(end)

        def apply_field_text(paragraph, sample):
            m = re.search(r'\d+', sample)
            if not m:
                paragraph.add_run(sample)
                return
            paragraph.add_run(sample[:m.start()])
            add_page_number_field(paragraph)
            paragraph.add_run(sample[m.end():])

        def overlap_ratio(a, b):
            ax0, ay0, ax1, ay1 = a
            bx0, by0, bx1, by1 = b
            ix0, iy0 = max(ax0, bx0), max(ay0, by0)
            ix1, iy1 = min(ax1, bx1), min(ay1, by1)
            if ix1 <= ix0 or iy1 <= iy0:
                return 0.0
            inter = (ix1 - ix0) * (iy1 - iy0)
            area_a = max(1.0, (ax1 - ax0) * (ay1 - ay0))
            return inter / area_a

        def detect_columns(items, page_x0, page_x1):
            # Finds vertical whitespace gutters shared by most lines, so multi-column
            # PDFs (magazines/papers) read top-to-bottom per column instead of
            # interleaving alternating lines from the left and right columns.
            page_width = page_x1 - page_x0
            if page_width <= 0 or len(items) < 6:
                return [(page_x0, page_x1)]
            bins = 80
            bin_w = page_width / bins
            covered = [False] * bins
            for it in items:
                x0, x1 = it["bbox"][0], it["bbox"][2]
                b0 = max(0, int((x0 - page_x0) / bin_w))
                b1 = min(bins - 1, int((x1 - page_x0) / bin_w))
                for i in range(b0, b1 + 1):
                    covered[i] = True
            gaps = []
            i = 0
            while i < bins:
                if not covered[i]:
                    j = i
                    while j < bins and not covered[j]:
                        j += 1
                    gaps.append((page_x0 + i * bin_w, page_x0 + j * bin_w, i, j - 1))
                    i = j
                else:
                    i += 1
            real_gaps = [g for g in gaps if (g[1] - g[0]) >= 12 and g[2] > 3 and g[3] < bins - 4]
            if not real_gaps or len(real_gaps) > 2:
                return [(page_x0, page_x1)]
            splits = sorted((g[0] + g[1]) / 2 for g in real_gaps)
            bands, prev = [], page_x0
            for s in splits:
                bands.append((prev, s))
                prev = s
            bands.append((prev, page_x1))
            counts = [sum(1 for it in items if bx0 <= (it["bbox"][0] + it["bbox"][2]) / 2 < bx1) for bx0, bx1 in bands]
            total = sum(counts)
            if total == 0 or min(counts) < max(2, total * 0.12):
                return [(page_x0, page_x1)]
            return bands

        def render_table(item, state):
            rows = item["rows"]
            max_cols = max(len(r) for r in rows)
            table = out_doc.add_table(rows=len(rows), cols=max_cols)
            table.style = "Table Grid"
            for r_idx, row in enumerate(rows):
                for c_idx in range(max_cols):
                    val = row[c_idx] if c_idx < len(row) and row[c_idx] is not None else ""
                    table.cell(r_idx, c_idx).text = str(val).strip()
            state["p"] = None
            state["last_y_bottom"] = -999
            state["last_font_size"] = -1

        def render_image(item, state):
            bbox = item["bbox"]
            width_in = min((bbox[2] - bbox[0]) / 72.0, usable_width_in)
            if width_in <= 0:
                width_in = usable_width_in
            try:
                out_doc.add_picture(BytesIO(item["data"]), width=Inches(width_in))
            except Exception:
                pass
            state["p"] = None
            state["last_y_bottom"] = -999
            state["last_font_size"] = -1

        def render_text(item, state):
            bbox = item["bbox"]
            y_top, y_bottom = bbox[1], bbox[3]
            font_size = item["font_size"]
            vertical_gap = y_top - state["last_y_bottom"]
            is_bullet = item["full_text"].startswith(('•', '◦', '-', '*'))

            start_new_para = True
            if state["p"] is not None:
                if is_bullet:
                    start_new_para = True
                elif y_top < (state["last_y_bottom"] - font_size * 0.3):
                    start_new_para = True
                elif vertical_gap < (font_size * 0.5) and abs(font_size - state["last_font_size"]) < 1.0:
                    start_new_para = False

            if start_new_para:
                state["p"] = out_doc.add_paragraph()
            else:
                state["p"].add_run(" ")

            for r in item["runs"]:
                run = state["p"].add_run(r["text"])
                raw_font = r["font"]
                flags = r["flags"]
                if (flags & 16) or "bold" in raw_font or "black" in raw_font:
                    run.bold = True
                if (flags & 2) or "italic" in raw_font:
                    run.italic = True
                clean_name = clean_font_name(raw_font)
                if clean_name and clean_name.lower() != "unknown":
                    run.font.name = clean_name
                if r["size"] > 0:
                    run.font.size = Pt(round(r["size"]))

            state["last_y_bottom"] = y_bottom
            state["last_font_size"] = font_size

        def render_band(band_items, page_x0, page_x1, state):
            if not band_items:
                return
            col_bands = detect_columns(band_items, page_x0, page_x1)
            if len(col_bands) <= 1:
                ordered = sorted(band_items, key=lambda x: (round(x["bbox"][1] / 5), x["bbox"][0]))
                for it in ordered:
                    (render_image if it["type"] == "image" else render_text)(it, state)
                return
            for bx0, bx1 in col_bands:
                col_items = [it for it in band_items if bx0 <= (it["bbox"][0] + it["bbox"][2]) / 2 < bx1]
                col_items.sort(key=lambda x: (round(x["bbox"][1] / 5), x["bbox"][0]))
                for it in col_items:
                    (render_image if it["type"] == "image" else render_text)(it, state)

        page_count = len(pdf)

        # ---- Pass 1: detect header/footer lines that repeat across most pages, so
        # they land in Word's real header/footer instead of the document body. ----
        top_counter, bottom_counter = Counter(), Counter()
        top_samples, bottom_samples = {}, {}
        for page in pdf:
            h = page.rect.height
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") != 0:
                    continue
                for l in b.get("lines", []):
                    txt = "".join(s.get("text", "") for s in l.get("spans", [])).strip()
                    if not txt:
                        continue
                    y0, y1 = l["bbox"][1], l["bbox"][3]
                    norm = re.sub(r'\d+', '#', txt)
                    if y1 < h * 0.10:
                        top_counter[norm] += 1
                        top_samples.setdefault(norm, txt)
                    elif y0 > h * 0.90:
                        bottom_counter[norm] += 1
                        bottom_samples.setdefault(norm, txt)

        repeat_threshold = max(2, int(page_count * 0.6))
        header_patterns = {p for p, c in top_counter.items() if c >= repeat_threshold} if page_count >= 3 else set()
        footer_patterns = {p for p, c in bottom_counter.items() if c >= repeat_threshold} if page_count >= 3 else set()

        if header_patterns:
            apply_field_text(section.header.paragraphs[0], top_samples[next(iter(header_patterns))])
        if footer_patterns:
            apply_field_text(section.footer.paragraphs[0], bottom_samples[next(iter(footer_patterns))])

        # ---- Pass 2: rebuild the body page by page (real page breaks, tables,
        # inline images and column-aware reading order) ----
        for page_index, page in enumerate(pdf):
            if page_index > 0:
                out_doc.add_page_break()

            h = page.rect.height
            page_x0, page_x1 = page.rect.x0, page.rect.x1

            table_items = []
            try:
                for t in page.find_tables().tables:
                    try:
                        rows = t.extract()
                    except Exception:
                        continue
                    if rows:
                        table_items.append({"type": "table", "bbox": list(t.bbox), "rows": rows})
            except Exception:
                pass

            image_items = []
            for img_info in page.get_images(full=True):
                xref = img_info[0]
                try:
                    rects = page.get_image_rects(xref) or []
                except Exception:
                    rects = []
                if not rects:
                    continue
                try:
                    pix = fitz.Pixmap(pdf, xref)
                    if pix.n - pix.alpha >= 4:
                        pix = fitz.Pixmap(fitz.csRGB, pix)
                    img_bytes = pix.tobytes("png")
                except Exception:
                    continue
                for rect in rects:
                    rbbox = [rect[0], rect[1], rect[2], rect[3]]
                    if (rbbox[2] - rbbox[0]) < 5 or (rbbox[3] - rbbox[1]) < 5:
                        continue
                    if any(overlap_ratio(rbbox, t["bbox"]) > 0.5 for t in table_items):
                        continue
                    image_items.append({"type": "image", "bbox": rbbox, "data": img_bytes})

            text_items = []
            for b in page.get_text("dict")["blocks"]:
                if b.get("type") != 0:
                    continue
                for l in b.get("lines", []):
                    spans = l.get("spans", [])
                    if not spans:
                        continue
                    bbox = l.get("bbox", [0, 0, 0, 0])
                    raw_text = "".join(s.get("text", "") for s in spans).strip()
                    if not raw_text:
                        continue
                    norm = re.sub(r'\d+', '#', raw_text)
                    in_header_zone = bbox[3] < h * 0.10
                    in_footer_zone = bbox[1] > h * 0.90
                    if (in_header_zone and norm in header_patterns) or (in_footer_zone and norm in footer_patterns):
                        continue
                    if any(overlap_ratio(bbox, t["bbox"]) > 0.5 for t in table_items):
                        continue

                    run_params = []
                    last_x = None
                    first_span = spans[0]
                    font_size = first_span.get("size", 11)

                    for s in spans:
                        text = s.get("text", "")
                        sbbox = s.get("bbox", [0, 0, 0, 0])

                        if last_x is not None:
                            gap = sbbox[0] - last_x
                            if gap > 2 and not text.startswith(" ") and not text.startswith(","):
                                run_params.append({"text": " ", "size": font_size, "font": first_span.get("font", ""), "flags": 0})

                        run_params.append({"text": text, "size": s.get("size", 11), "font": s.get("font", "").lower(), "flags": s.get("flags", 0)})
                        last_x = sbbox[2]

                    full_text = "".join(r["text"] for r in run_params).strip()
                    if not full_text:
                        continue
                    text_items.append({"type": "text", "bbox": bbox, "runs": run_params, "full_text": full_text, "font_size": font_size})

            content_width = max((it["bbox"][2] for it in (text_items + image_items)), default=page_x1) - page_x0
            flowable = sorted(text_items + image_items + table_items, key=lambda it: (round(it["bbox"][1] / 5), it["bbox"][0]))

            state = {"p": None, "last_y_bottom": -999, "last_font_size": -1}
            band = []
            for it in flowable:
                is_break = it["type"] == "table" or (it["type"] == "image" and (it["bbox"][2] - it["bbox"][0]) >= 0.7 * max(content_width, 1))
                if is_break:
                    render_band(band, page_x0, page_x1, state)
                    band = []
                    (render_table if it["type"] == "table" else render_image)(it, state)
                else:
                    band.append(it)
            render_band(band, page_x0, page_x1, state)

        out_doc.save(docx_path)
        
    # 1. Primary conversion using pdf2docx (fast, good enough for most straightforward PDFs)
    cv = Converter(temp_pdf_path)
    cv.convert(temp_docx_path, start=0, end=None)
    cv.close()

    used_method = "layout-mode"
    if has_spacing_issue(temp_docx_path):
        if fix_glued_words_in_docx(temp_docx_path):
            used_method = "layout-mode-corrected"
        else:
            convert_with_correct_spacing(temp_pdf_path, temp_docx_path)
            used_method = "text-mode"
    return used_method


@app.post("/convert/pdf-to-word")
async def convert_pdf_to_word(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    # Create temporary files for the input PDF and output DOCX
    fd_pdf, temp_pdf_path = tempfile.mkstemp(suffix=".pdf")
    fd_docx, temp_docx_path = tempfile.mkstemp(suffix=".docx")

    os.close(fd_pdf)
    os.close(fd_docx)

    try:
        used_method = await run_in_threadpool(_convert_pdf_to_word_sync, file.file, temp_pdf_path, temp_docx_path)

        background_tasks.add_task(remove_files, [temp_pdf_path, temp_docx_path])
        return FileResponse(
            path=temp_docx_path,
            filename=f"converted_{file.filename}.docx",
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            headers={"X-Conversion-Method": used_method}
        )
    except Exception as e:
        remove_files([temp_pdf_path, temp_docx_path])
        raise HTTPException(status_code=500, detail=str(e))

def _pdf_cell_to_excel(raw):
    """Turn a scraped table cell into a native Excel type (+ number format) where it's
    unambiguously safe to do so, so numbers land as real numbers (sortable/summable,
    right-aligned) instead of importing as text - the single biggest visible gap between
    a scraped-looking sheet and one a person typed by hand. Returns (value, number_format)."""
    if raw is None:
        return None, None
    text = str(raw).strip()
    if text == "":
        return None, None
    # Leading-zero strings are almost always identifiers (zip codes, IDs), not numbers -
    # converting "00123" to 123 would silently corrupt the data.
    if re.fullmatch(r"0\d+", text):
        return text, None
    cleaned = text
    is_negative = cleaned.startswith("(") and cleaned.endswith(")")
    if is_negative:
        cleaned = cleaned[1:-1]
    cleaned = cleaned.replace(",", "")
    is_percent = cleaned.endswith("%")
    if is_percent:
        cleaned = cleaned[:-1]
    currency_match = re.match(r"^([$€£])", cleaned)
    if currency_match:
        cleaned = cleaned[1:]
    if not re.fullmatch(r"-?\d+(\.\d+)?", cleaned or ""):
        return text, None
    num = float(cleaned)
    if is_negative:
        num = -num
    if is_percent:
        return num / 100, "0.00%"
    if currency_match:
        return num, f'{currency_match.group(1)}#,##0.00'
    if num == int(num) and "." not in cleaned:
        return int(num), None
    return num, None


def _convert_pdf_to_excel_sync(file_obj, temp_pdf_path: str, temp_xlsx_path: str) -> None:
    """All the blocking work for pdf-to-excel: writing the upload to disk and the
    pdfplumber-based table extraction below (synchronous, CPU/IO-bound). Must be called
    via run_in_threadpool - otherwise a large/complex PDF blocks every other request this
    service is handling for however long extraction takes."""
    with open(temp_pdf_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    import pdfplumber

    wb = Workbook()
    ws = wb.active
    header_font = Font(bold=True)
    max_col_widths = {}

    def write_row(row_idx, values, bold=False):
        for col_idx, raw in enumerate(values, start=1):
            value, number_format = _pdf_cell_to_excel(raw)
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            if number_format:
                cell.number_format = number_format
            if bold:
                cell.font = header_font
            width = len(str(raw)) if raw is not None else 0
            if width > max_col_widths.get(col_idx, 0):
                max_col_widths[col_idx] = width

    current_row = 1
    prev_num_cols = None
    prev_header = None

    with pdfplumber.open(temp_pdf_path) as pdf:
        for page in pdf.pages:
            tables = [t for t in page.extract_tables() if t]

            if not tables:
                # Robust unruled / borderless table column-region extraction:
                # Detects vertical column channels across the page by analyzing word coordinates,
                # intra-phrase spacing, and alignments. Ensures that right-aligned numbers
                # (e.g. Age: 32) and adjacent date columns (e.g. 15/10/2017) or names (First/Last)
                # never merge into one cell.
                extracted_rows = []
                words = page.extract_words(x_tolerance=1, y_tolerance=1)
                if words:
                    lines = []
                    current_line = []
                    for w in sorted(words, key=lambda x: (x['top'], x['x0'])):
                        if not current_line:
                            current_line.append(w)
                        else:
                            if abs(w['top'] - current_line[0]['top']) <= 4.0:
                                current_line.append(w)
                            else:
                                lines.append(sorted(current_line, key=lambda x: x['x0']))
                                current_line = [w]
                    if current_line:
                        lines.append(sorted(current_line, key=lambda x: x['x0']))

                    table_lines = []
                    for line in lines:
                        if len(line) == 1 and line[0]['text'] in ['Sheet1', 'Page']:
                            continue
                        if len(line) <= 2 and 'Page' in [w['text'] for w in line]:
                            continue
                        table_lines.append(line)

                    if table_lines:
                        line_phrases = []
                        for line in table_lines:
                            phrases = []
                            cur_phrase = []
                            for w in line:
                                if not cur_phrase:
                                    cur_phrase.append(w)
                                else:
                                    prev_w = cur_phrase[-1]
                                    gap = w['x0'] - prev_w['x1']
                                    if gap <= 3.2:
                                        cur_phrase.append(w)
                                    else:
                                        phrases.append(cur_phrase)
                                        cur_phrase = [w]
                            if cur_phrase:
                                phrases.append(cur_phrase)
                            line_phrases.append(phrases)

                        all_intervals = []
                        for phrases in line_phrases:
                            for p in phrases:
                                all_intervals.append((p[0]['x0'], p[-1]['x1']))

                        page_w = int(page.width) + 1
                        coverage = [0] * page_w
                        for x0, x1 in all_intervals:
                            for x in range(max(0, int(x0 + 0.5)), min(page_w, int(x1 + 0.5))):
                                coverage[x] += 1

                        active_regions = []
                        in_active = False
                        reg_start = 0
                        for x in range(page_w):
                            if coverage[x] > 0 and not in_active:
                                in_active = True
                                reg_start = x
                            elif coverage[x] == 0 and in_active:
                                in_active = False
                                active_regions.append([reg_start, x])
                        if in_active:
                            active_regions.append([reg_start, page_w])

                        # Merge complementary sub-tracks (e.g. left-aligned header + right-aligned numbers)
                        merged_regions = []
                        i = 0
                        while i < len(active_regions):
                            cur_start, cur_end = active_regions[i]
                            if i + 1 < len(active_regions):
                                next_start, next_end = active_regions[i + 1]
                                overlap = False
                                for phrases in line_phrases:
                                    has_cur = any(cur_start - 2 <= (p[0]['x0']+p[-1]['x1'])/2 <= cur_end + 2 for p in phrases)
                                    has_next = any(next_start - 2 <= (p[0]['x0']+p[-1]['x1'])/2 <= next_end + 2 for p in phrases)
                                    if has_cur and has_next:
                                        overlap = True
                                        break
                                if not overlap and (next_end - cur_start) < 100:
                                    merged_regions.append([cur_start, next_end])
                                    i += 2
                                    continue
                            merged_regions.append([cur_start, cur_end])
                            i += 1

                        for phrases in line_phrases:
                            row_cells = [[] for _ in merged_regions]
                            for p in phrases:
                                p_text = " ".join([w['text'] for w in p])
                                p_mid = (p[0]['x0'] + p[-1]['x1']) / 2.0
                                matched_idx = -1
                                for idx, (t_start, t_end) in enumerate(merged_regions):
                                    if t_start - 2 <= p_mid <= t_end + 2:
                                        matched_idx = idx
                                        break
                                if matched_idx == -1:
                                    distances = [abs(p_mid - (t_start + t_end)/2.0) for t_start, t_end in merged_regions]
                                    matched_idx = distances.index(min(distances))
                                row_cells[matched_idx].append(p_text)

                            row = [" ".join(c) if c else "" for c in row_cells]
                            while row and row[-1] == "":
                                row.pop()
                            if any(row):
                                extracted_rows.append(row)

                if extracted_rows:
                    if prev_num_cols is not None:
                        current_row += 1
                    for i, row in enumerate(extracted_rows):
                        is_header = (i == 0 and len(row) > 1)
                        write_row(current_row, row, bold=is_header)
                        current_row += 1
                    prev_num_cols = len(extracted_rows[0]) if extracted_rows else None
                    prev_header = None
                continue

            for table in tables:
                raw_rows = [["" if c is None else str(c).strip() for c in row] for row in table]
                num_cols = len(raw_rows[0]) if raw_rows else 0

                # A table that resumes on the next page re-extracts with the same
                # column count and a repeated header row - treat it as a continuation
                # of the previous block (no blank separator, no duplicate header)
                # rather than a brand-new table.
                is_continuation = (
                    prev_num_cols == num_cols
                    and prev_header is not None
                    and raw_rows
                    and raw_rows[0] == prev_header
                )

                rows_to_write = raw_rows[1:] if is_continuation else raw_rows
                if not is_continuation and prev_num_cols is not None:
                    current_row += 1  # blank row between genuinely distinct tables

                for i, row in enumerate(rows_to_write):
                    is_header_row = not is_continuation and i == 0
                    write_row(current_row, row, bold=is_header_row)
                    current_row += 1

                prev_num_cols = num_cols
                prev_header = raw_rows[0] if raw_rows else None

    for col_idx, width in max_col_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 8), 60)

    wb.save(temp_xlsx_path)


@app.post("/convert/pdf-to-excel")
async def convert_pdf_to_excel(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    fd_pdf, temp_pdf_path = tempfile.mkstemp(suffix=".pdf")
    fd_xlsx, temp_xlsx_path = tempfile.mkstemp(suffix=".xlsx")

    os.close(fd_pdf)
    os.close(fd_xlsx)

    try:
        await run_in_threadpool(_convert_pdf_to_excel_sync, file.file, temp_pdf_path, temp_xlsx_path)

        background_tasks.add_task(remove_files, [temp_pdf_path, temp_xlsx_path])
        return FileResponse(path=temp_xlsx_path, filename=f"converted_{file.filename}.xlsx", media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    except Exception as e:
        remove_files([temp_pdf_path, temp_xlsx_path])
        raise HTTPException(status_code=500, detail=str(e))

def _convert_pdf_to_ppt_sync(file_obj, temp_pdf_path: str, temp_pptx_path: str) -> None:
    """All the blocking work for pdf-to-ppt: writing the upload to disk and rebuilding each
    page as real, editable PowerPoint content (text boxes, native tables, and images placed
    at their original PDF positions) rather than a flat screenshot. Synchronous/CPU-bound -
    must be called via run_in_threadpool, otherwise a large/multi-page PDF blocks every other
    request this service is handling for however long rebuilding takes.

    Each page is reconstructed independently (unlike the pdf-to-word text-mode path, a slide
    is an absolute canvas, so there's no paragraph-flow/column-reading-order to work out -
    every element is just placed at its own bbox). PPTX requires one slide size for the whole
    deck, so pages smaller than the largest page are centered on it instead of stretched - a
    page is never resized to fit, so nothing on it is ever distorted. A page whose content is
    pure vector art (no extractable text or raster images - e.g. some diagrams/illustrations)
    falls back to a full-page screenshot of just that page, so no slide ever ends up blank."""
    from io import BytesIO
    from pptx.util import Pt
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.dml.color import RGBColor

    with open(temp_pdf_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    doc = open_pdf_or_raise(temp_pdf_path)
    prs = Presentation()
    blank_slide_layout = prs.slide_layouts[6] # blank layout

    if len(doc) == 0:
        prs.save(temp_pptx_path)
        return

    def clean_font_name(f_str):
        name = (f_str or "").split('+')[-1]
        name = re.sub(r'-(Bold|Italic|Regular|Medium|SemiBold|Light).*', '', name, flags=re.IGNORECASE)
        name = re.sub(r'(MT|PSMT|MS|PS)$', '', name, flags=re.IGNORECASE)
        name = re.sub(r'([a-z])([A-Z])', r'\1 \2', name)
        return name.strip()

    def overlap_ratio(a, b):
        ax0, ay0, ax1, ay1 = a
        bx0, by0, bx1, by1 = b
        ix0, iy0 = max(ax0, bx0), max(ay0, by0)
        ix1, iy1 = min(ax1, bx1), min(ay1, by1)
        if ix1 <= ix0 or iy1 <= iy0:
            return 0.0
        inter = (ix1 - ix0) * (iy1 - iy0)
        area_a = max(1.0, (ax1 - ax0) * (ay1 - ay0))
        return inter / area_a

    def add_table(slide, item, offset_x, offset_y):
        bx0, by0, bx1, by1 = item["bbox"]
        rows = item["rows"]
        num_cols = max((len(r) for r in rows), default=0)
        if not rows or num_cols == 0:
            return
        width = max(bx1 - bx0, 10)
        height = max(by1 - by0, 10 * len(rows))
        graphic_frame = slide.shapes.add_table(
            len(rows), num_cols, Pt(bx0 + offset_x), Pt(by0 + offset_y), Pt(width), Pt(height)
        )
        table = graphic_frame.table
        for r_idx, row in enumerate(rows):
            for c_idx in range(num_cols):
                val = row[c_idx] if c_idx < len(row) and row[c_idx] is not None else ""
                table.cell(r_idx, c_idx).text = str(val).strip()

    def add_image(slide, item, offset_x, offset_y):
        bx0, by0, bx1, by1 = item["bbox"]
        width = max(bx1 - bx0, 1)
        height = max(by1 - by0, 1)
        try:
            slide.shapes.add_picture(
                BytesIO(item["data"]), Pt(bx0 + offset_x), Pt(by0 + offset_y), width=Pt(width), height=Pt(height)
            )
        except Exception:
            pass

    def add_text_line(slide, item, offset_x, offset_y):
        bx0, by0, bx1, by1 = item["bbox"]
        font_size = item["font_size"]
        box_width = max(bx1 - bx0, 1) + 4 # small margin so a run is never forced to wrap
        box_height = max(by1 - by0, font_size * 1.3)

        textbox = slide.shapes.add_textbox(Pt(bx0 + offset_x), Pt(by0 + offset_y), Pt(box_width), Pt(box_height))
        tf = textbox.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        paragraph = tf.paragraphs[0]

        for r in item["runs"]:
            run = paragraph.add_run()
            run.text = r["text"]
            if r["size"] > 0:
                run.font.size = Pt(round(r["size"]))
            raw_font = r["font"]
            flags = r["flags"]
            if (flags & 16) or "bold" in raw_font or "black" in raw_font:
                run.font.bold = True
            if (flags & 2) or "italic" in raw_font:
                run.font.italic = True
            clean_name = clean_font_name(r["font"])
            if clean_name and clean_name.lower() != "unknown":
                run.font.name = clean_name
            color = r.get("color")
            if isinstance(color, int):
                try:
                    run.font.color.rgb = RGBColor((color >> 16) & 255, (color >> 8) & 255, color & 255)
                except Exception:
                    pass

    def add_fallback_screenshot(slide, page, offset_x, offset_y, page_w, page_h):
        # Safety net for pages with no extractable text/images (pure vector art, e.g. some
        # diagrams) - without this, such a page would render as a silently blank slide.
        pix = page.get_pixmap(dpi=150)
        img_data = pix.tobytes("png")
        try:
            slide.shapes.add_picture(BytesIO(img_data), Pt(offset_x), Pt(offset_y), width=Pt(page_w), height=Pt(page_h))
        except Exception:
            pass

    # PPTX requires one slide size for the whole deck - use the largest page in each
    # dimension so no page's content is ever scaled down, then center every page on it.
    max_w = max(page.rect.width for page in doc)
    max_h = max(page.rect.height for page in doc)
    prs.slide_width = Pt(max_w)
    prs.slide_height = Pt(max_h)

    for page in doc:
        slide = prs.slides.add_slide(blank_slide_layout)
        page_w, page_h = page.rect.width, page.rect.height
        offset_x = (max_w - page_w) / 2.0
        offset_y = (max_h - page_h) / 2.0

        table_items = []
        try:
            for t in page.find_tables().tables:
                try:
                    rows = t.extract()
                except Exception:
                    continue
                if rows:
                    table_items.append({"bbox": list(t.bbox), "rows": rows})
        except Exception:
            pass

        image_items = []
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            try:
                rects = page.get_image_rects(xref) or []
            except Exception:
                rects = []
            if not rects:
                continue
            try:
                pix = fitz.Pixmap(doc, xref)
                if pix.n - pix.alpha >= 4:
                    pix = fitz.Pixmap(fitz.csRGB, pix)
                img_bytes = pix.tobytes("png")
            except Exception:
                continue
            for rect in rects:
                rbbox = [rect[0], rect[1], rect[2], rect[3]]
                if (rbbox[2] - rbbox[0]) < 2 or (rbbox[3] - rbbox[1]) < 2:
                    continue
                if any(overlap_ratio(rbbox, t["bbox"]) > 0.5 for t in table_items):
                    continue
                image_items.append({"bbox": rbbox, "data": img_bytes})

        text_items = []
        for b in page.get_text("dict")["blocks"]:
            if b.get("type") != 0:
                continue
            for l in b.get("lines", []):
                spans = l.get("spans", [])
                if not spans:
                    continue
                bbox = l.get("bbox", [0, 0, 0, 0])
                raw_text = "".join(s.get("text", "") for s in spans).strip()
                if not raw_text:
                    continue
                if any(overlap_ratio(bbox, t["bbox"]) > 0.5 for t in table_items):
                    continue

                run_params = []
                last_x = None
                first_span = spans[0]
                font_size = first_span.get("size", 11) or 11

                for s in spans:
                    text = s.get("text", "")
                    sbbox = s.get("bbox", [0, 0, 0, 0])
                    if last_x is not None:
                        gap = sbbox[0] - last_x
                        if gap > 2 and not text.startswith(" ") and not text.startswith(","):
                            run_params.append({"text": " ", "size": font_size, "font": first_span.get("font", ""), "flags": 0, "color": first_span.get("color")})
                    run_params.append({
                        "text": text,
                        "size": s.get("size", 11) or 11,
                        "font": (s.get("font") or "").lower(),
                        "flags": s.get("flags", 0),
                        "color": s.get("color"),
                    })
                    last_x = sbbox[2]

                text_items.append({"bbox": bbox, "runs": run_params, "font_size": font_size})

        if not table_items and not image_items and not text_items:
            add_fallback_screenshot(slide, page, offset_x, offset_y, page_w, page_h)
            continue

        for t in table_items:
            add_table(slide, t, offset_x, offset_y)
        for img in image_items:
            add_image(slide, img, offset_x, offset_y)
        for line in text_items:
            add_text_line(slide, line, offset_x, offset_y)

    prs.save(temp_pptx_path)


@app.post("/convert/pdf-to-ppt")
async def convert_pdf_to_ppt(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    fd_pdf, temp_pdf_path = tempfile.mkstemp(suffix=".pdf")
    fd_pptx, temp_pptx_path = tempfile.mkstemp(suffix=".pptx")

    os.close(fd_pdf)
    os.close(fd_pptx)

    try:
        await run_in_threadpool(_convert_pdf_to_ppt_sync, file.file, temp_pdf_path, temp_pptx_path)

        background_tasks.add_task(remove_files, [temp_pdf_path, temp_pptx_path])
        return FileResponse(path=temp_pptx_path, filename=f"converted_{file.filename}.pptx", media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation")
    except InvalidFileError as e:
        remove_files([temp_pdf_path, temp_pptx_path])
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        remove_files([temp_pdf_path, temp_pptx_path])
        print(f"[pdf-to-ppt] conversion failed: {e}")
        raise HTTPException(
            status_code=500,
            detail="We couldn't convert this file due to an unexpected server error. Please try again shortly.",
        )

def _convert_pdf_to_images_sync(file_obj, temp_pdf_path: str, temp_zip_path: str) -> None:
    """All the blocking work for pdf-to-images: writing the upload to disk and rendering each
    page to a JPEG (PyMuPDF, synchronous/CPU-bound). Must be called via run_in_threadpool -
    otherwise a large/multi-page PDF blocks every other request this service is handling for
    however long rendering takes."""
    with open(temp_pdf_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    doc = fitz.open(temp_pdf_path)

    with zipfile.ZipFile(temp_zip_path, 'w') as zipf:
        for i, page in enumerate(doc):
            pix = page.get_pixmap(dpi=150)
            img_data = pix.tobytes("jpeg")
            zipf.writestr(f"page_{i+1}.jpg", img_data)


@app.post("/convert/pdf-to-images")
async def convert_pdf_to_images(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    fd_pdf, temp_pdf_path = tempfile.mkstemp(suffix=".pdf")
    fd_zip, temp_zip_path = tempfile.mkstemp(suffix=".zip")

    os.close(fd_pdf)
    os.close(fd_zip)

    try:
        await run_in_threadpool(_convert_pdf_to_images_sync, file.file, temp_pdf_path, temp_zip_path)

        background_tasks.add_task(remove_files, [temp_pdf_path, temp_zip_path])
        return FileResponse(path=temp_zip_path, filename=f"converted_images.zip", media_type="application/zip")
    except Exception as e:
        remove_files([temp_pdf_path, temp_zip_path])
        raise HTTPException(status_code=500, detail=str(e))

@app.post("/convert/text-to-speech")
async def convert_text_to_speech(background_tasks: BackgroundTasks, request: TTSRequest):
    fd_mp3, temp_mp3_path = tempfile.mkstemp(suffix=".mp3")
    os.close(fd_mp3)

    try:
        # Edge-tts Communicate object
        communicate = edge_tts.Communicate(request.text, request.voice, rate=request.speed)
        
        # Generate and save audio to temp file
        await communicate.save(temp_mp3_path)
        
        background_tasks.add_task(remove_files, [temp_mp3_path])
        return FileResponse(path=temp_mp3_path, filename="speech.mp3", media_type="audio/mpeg")
    except Exception as e:
        remove_files([temp_mp3_path])
        print(f"[text-to-speech] edge-tts failed: {e}")
        # Must raise (not return a dict) so the response carries a non-200 status -
        # the Next.js caller requests this as arraybuffer and can't tell a JSON
        # error body from real audio bytes otherwise, so a 200 here silently
        # saves the error message as if it were a working MP3.
        raise HTTPException(status_code=500, detail=str(e))

def _run_whisper_transcription(audio_path: str, transcribe_kwargs: dict) -> tuple[str, str]:
    """Runs the actual (CPU-bound, synchronous) faster-whisper inference. Must be
    called via run_in_threadpool - transcribe() returns a lazily-evaluated generator,
    so the real work happens while iterating segments below, not in the transcribe()
    call itself. Called directly on the event loop this would block every other
    request the service is handling (TTS, PDF conversions, other users' STT) for the
    entire duration of the transcription."""
    segments, info = whisper_model.transcribe(audio_path, **transcribe_kwargs)
    full_text = ""
    for segment in segments:
        full_text += segment.text + " "
    return full_text.strip(), info.language


@app.post("/convert/speech-to-text")
async def convert_speech_to_text(file: UploadFile = File(...), language: str = Form(None)):
    # Save the uploaded file temporarily
    suffix = os.path.splitext(file.filename or "")[1]
    fd_audio, temp_audio_path = tempfile.mkstemp(suffix=suffix)
    os.close(fd_audio)

    try:
        with open(temp_audio_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        # Auto-detection on the small "base" model can misfire for lower-resource
        # languages (e.g. Urdu getting misdetected as Chinese) - when the caller
        # knows the spoken language, passing it explicitly skips that guess and
        # is noticeably more accurate than relying on auto-detect.
        transcribe_kwargs = {"beam_size": 5}
        if language:
            transcribe_kwargs["language"] = language

        full_text, detected_language = await run_in_threadpool(
            _run_whisper_transcription, temp_audio_path, transcribe_kwargs
        )

        return {"text": full_text, "language": detected_language}
    except Exception as e:
        return {"error": str(e)}
    finally:
        if os.path.exists(temp_audio_path):
            try:
                os.remove(temp_audio_path)
            except:
                pass

@app.post("/convert/word-to-pdf")
async def convert_word_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    return await office_to_pdf_response(background_tasks, file)


@app.post("/convert/excel-to-pdf")
async def convert_excel_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    return await office_to_pdf_response(background_tasks, file)


@app.post("/convert/ppt-to-pdf")
async def convert_ppt_to_pdf(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    return await office_to_pdf_response(background_tasks, file)


def _compress_pdf_sync(file_obj, input_path: str, output_path: str, pdf_settings: str) -> tuple[int, int]:
    """All the blocking work for compress-pdf: writing the upload to disk and running
    Ghostscript. Must be called via run_in_threadpool - Ghostscript can take up to its 100s
    subprocess timeout, and calling it directly on the event loop would freeze every other
    request this service is handling for that entire duration."""
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)
    before_size = os.path.getsize(input_path)

    result = subprocess.run(
        [
            GHOSTSCRIPT_BIN,
            "-sDEVICE=pdfwrite",
            "-dCompatibilityLevel=1.4",
            f"-dPDFSETTINGS={pdf_settings}",
            "-dNOPAUSE",
            "-dQUIET",
            "-dBATCH",
            f"-sOutputFile={output_path}",
            input_path,
        ],
        capture_output=True,
        timeout=100,
    )
    if result.returncode != 0 or not os.path.exists(output_path):
        stderr = result.stderr.decode(errors="ignore") if result.stderr else "unknown error"
        raise RuntimeError(f"Ghostscript compression failed: {stderr}")
    after_size = os.path.getsize(output_path)
    return before_size, after_size


@app.post("/convert/compress-pdf")
async def convert_compress_pdf(
    background_tasks: BackgroundTasks, file: UploadFile = File(...), level: str = Form("recommended")
):
    pdf_settings = {"low": "/printer", "recommended": "/ebook", "extreme": "/screen"}.get(level, "/ebook")
    work_dir = tempfile.mkdtemp()
    input_path = os.path.join(work_dir, "input.pdf")
    output_path = os.path.join(work_dir, "output.pdf")

    try:
        before_size, after_size = await run_in_threadpool(_compress_pdf_sync, file.file, input_path, output_path, pdf_settings)

        background_tasks.add_task(shutil.rmtree, work_dir, True)
        return FileResponse(
            path=output_path,
            filename="Compressed.pdf",
            media_type="application/pdf",
            headers={"X-Before-Size": str(before_size), "X-After-Size": str(after_size)},
        )
    except Exception as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise HTTPException(status_code=500, detail=str(e))


def _pdf_thumbnail_sync(file_obj, temp_pdf_path: str, temp_png_path: str) -> None:
    """The blocking work for pdf-thumbnail (writing the upload to disk and rendering page 1
    via PyMuPDF). Cheap per-call, but still run via run_in_threadpool for consistency with
    every other endpoint here, and so a burst of thumbnail requests can't add up to a
    noticeable stall on the event loop."""
    with open(temp_pdf_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    doc = fitz.open(temp_pdf_path)
    if len(doc) == 0:
        raise RuntimeError("PDF has no pages")
    doc[0].get_pixmap(dpi=72).save(temp_png_path)


@app.post("/convert/pdf-thumbnail")
async def convert_pdf_thumbnail(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    fd_pdf, temp_pdf_path = tempfile.mkstemp(suffix=".pdf")
    fd_png, temp_png_path = tempfile.mkstemp(suffix=".png")
    os.close(fd_pdf)
    os.close(fd_png)

    try:
        await run_in_threadpool(_pdf_thumbnail_sync, file.file, temp_pdf_path, temp_png_path)

        background_tasks.add_task(remove_files, [temp_pdf_path, temp_png_path])
        return FileResponse(path=temp_png_path, media_type="image/png")
    except Exception as e:
        remove_files([temp_pdf_path, temp_png_path])
        raise HTTPException(status_code=500, detail=str(e))


def _office_thumbnail_sync(file_obj, input_path: str, work_dir: str) -> str:
    """All the blocking work for office-thumbnail: writing the upload to disk, converting it
    to PDF via LibreOffice, then rendering page 1. Must be called via run_in_threadpool -
    LibreOffice alone can take up to its 100s subprocess timeout, and calling it directly on
    the event loop would freeze every other request this service is handling for that
    entire duration."""
    with open(input_path, "wb") as buffer:
        shutil.copyfileobj(file_obj, buffer)

    pdf_path = convert_with_soffice(input_path, work_dir)
    doc = fitz.open(pdf_path)
    if len(doc) == 0:
        raise RuntimeError("Converted document has no pages")
    png_path = os.path.join(work_dir, "thumb.png")
    doc[0].get_pixmap(dpi=72).save(png_path)
    return png_path


@app.post("/convert/office-thumbnail")
async def convert_office_thumbnail(background_tasks: BackgroundTasks, file: UploadFile = File(...)):
    ext = os.path.splitext(file.filename or "")[1] or ".tmp"
    work_dir = tempfile.mkdtemp()
    input_path = os.path.join(work_dir, f"input{ext}")

    try:
        png_path = await run_in_threadpool(_office_thumbnail_sync, file.file, input_path, work_dir)

        background_tasks.add_task(shutil.rmtree, work_dir, True)
        return FileResponse(path=png_path, media_type="image/png")
    except Exception as e:
        shutil.rmtree(work_dir, ignore_errors=True)
        raise HTTPException(status_code=500, detail=str(e))
